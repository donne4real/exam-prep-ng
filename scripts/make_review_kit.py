#!/usr/bin/env python3
"""Generate the teacher review kit for heuristic topic tags.

Output: review/topic-tags-review.xlsx
  - 'How to review' sheet with instructions and the full topic list
  - 'Review' sheet: a stratified random sample of tagged questions with
    blank columns for a teacher to mark tags right/wrong and suggest the
    correct topic
  - 'Coverage' sheet: tagging statistics per subject

Sampling is seeded, so re-running on the same bank produces the same kit.

Run:  python scripts/make_review_kit.py   (after scripts/build_bank.py)
"""

from __future__ import annotations

import json
import random
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).parent.parent
INDEX = ROOT / "public" / "data" / "index.json"
BANKS = ROOT / "public" / "data"
OUT = ROOT / "review" / "topic-tags-review.xlsx"

PER_TOPIC_SAMPLE = 3   # questions to sample per topic
UNTAGGED_SAMPLE = 5    # untagged questions to sample per subject
MAX_PER_SUBJECT = 30   # cap rows per subject so the kit stays reviewable

HEADER_FILL = PatternFill("solid", fgColor="008751")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SUBJECT_FILL = PatternFill("solid", fgColor="E6F4EE")
WRAP = Alignment(wrap_text=True, vertical="top")


def load_bank() -> list[dict]:
    with open(INDEX, encoding="utf-8") as fh:
        index = json.load(fh)
    questions: list[dict] = []
    for subject in index.get("subjects", []):
        path = BANKS / subject["file"]
        with open(path, encoding="utf-8") as fh:
            bank = json.load(fh)
        questions.extend(bank.get("questions", []))
    return questions


def build_workbook(questions: list[dict]) -> Workbook:
    rng = random.Random(42)
    wb = Workbook()

    # ---- Sheet 1: instructions ----
    ws = wb.active
    ws.title = "How to review"
    topics_by_subject: dict[tuple[str, str], Counter] = defaultdict(Counter)
    for q in questions:
        key = (q["exam"], q["subject"])
        topics_by_subject[key][q.get("topic") or "(untagged)"] += 1
    lines = [
        ("Topic tag review — ExamPrep NG question bank", True),
        ("", False),
        ("Every question in the app is automatically tagged with a syllabus topic", False),
        ("using keyword rules. These tags drive the 'weak topics' dashboard, so", False),
        ("wrong tags misdirect revision. This kit lets a teacher check a sample.", False),
        ("", False),
        ("How to review:", True),
        ("1. Go to the 'Review' sheet — questions are grouped by exam + subject.", False),
        ("2. For each row, read the question and its assigned topic.", False),
        ("3. If the topic is acceptable, leave 'Tag OK?' blank or write Y.", False),
        ("4. If it is wrong, write N and put the right topic in 'Correct topic'.", False),
        ("   Use the topic list at the bottom of this sheet (per subject) so the", False),
        ("   fix can be applied programmatically, or write free text for anything", False),
        ("   missing from the list.", False),
        ("5. Send the completed file back; the fixes go into scripts/build_bank.py.", False),
        ("", False),
        ("Notes:", True),
        ("- 'Tag OK?' accepts a stricter reading: the tag should be the syllabus", False),
        ("  topic a teacher would file the question under, not just related.", False),
        ("- The sample is random but seeded — the same questions can be discussed", False),
        ("  consistently between reviewers.", False),
    ]
    for row, (text, bold) in enumerate(lines, start=1):
        cell = ws.cell(row=row, column=1, value=text)
        cell.font = Font(bold=bold, size=14 if row == 1 else 11)
    ws.column_dimensions["A"].width = 100

    start = len(lines) + 2
    ws.cell(row=start, column=1, value="Topic list per subject (with question counts):").font = Font(bold=True)
    r = start + 1
    for (exam, subject), counter in sorted(topics_by_subject.items()):
        ws.cell(row=r, column=1, value=f"{exam} — {subject}").font = Font(bold=True)
        r += 1
        for topic, n in counter.most_common():
            ws.cell(row=r, column=2, value=f"{topic} ({n})")
            r += 1
        r += 1

    # ---- Sheet 2: review sample ----
    ws = wb.create_sheet("Review")
    headers = [
        "Exam", "Subject", "Question ID", "Year", "Question", "Options",
        "Correct", "Assigned topic", "Tag OK? (Y/N)", "Correct topic (if N)", "Notes",
    ]
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill, cell.font, cell.alignment = HEADER_FILL, HEADER_FONT, WRAP

    by_subject: dict[tuple[str, str], list[dict]] = defaultdict(list)
    for q in questions:
        by_subject[(q["exam"], q["subject"])].append(q)

    row = 2
    for (exam, subject), qs in sorted(by_subject.items()):
        by_topic: dict[str | None, list[dict]] = defaultdict(list)
        for q in qs:
            by_topic[q.get("topic")].append(q)
        sample: list[dict] = []
        for topic, topic_qs in sorted(by_topic.items(), key=lambda kv: (kv[0] is None, kv[0] or "")):
            if topic is None:
                continue
            sample.extend(rng.sample(topic_qs, min(PER_TOPIC_SAMPLE, len(topic_qs))))
        if None in by_topic:
            sample.extend(rng.sample(by_topic[None], min(UNTAGGED_SAMPLE, len(by_topic[None]))))
        rng.shuffle(sample)
        sample = sample[:MAX_PER_SUBJECT]

        # subject banner row
        cell = ws.cell(row=row, column=1, value=f"{exam} — {subject} ({len(sample)} sampled of {len(qs)})")
        for col in range(1, len(headers) + 1):
            ws.cell(row=row, column=col).fill = SUBJECT_FILL
        cell.font = Font(bold=True)
        row += 1

        for q in sample:
            options = " | ".join(f"{o['id'].upper()}. {o['text']}" for o in q["options"])
            correct = next(
                (o["text"] for o in q["options"] if o["id"] == q["correctOptionId"]), ""
            )
            values = [
                exam, subject, q["id"], q.get("year", ""),
                q["prompt"][:250], options[:300], correct,
                q.get("topic") or "(untagged)", "", "", "",
            ]
            for col, v in enumerate(values, start=1):
                c = ws.cell(row=row, column=col, value=v)
                c.alignment = WRAP
            row += 1
        row += 1  # blank line between subjects

    widths = [7, 18, 26, 7, 60, 50, 22, 22, 13, 22, 20]
    for col, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "A2"

    # ---- Sheet 3: coverage stats ----
    ws = wb.create_sheet("Coverage")
    for col, h in enumerate(
        ["Exam", "Subject", "Questions", "Tagged", "Untagged", "Tagged %", "Distinct topics"], start=1
    ):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    r = 2
    for (exam, subject), counter in sorted(topics_by_subject.items()):
        total = sum(counter.values())
        untagged = counter.get("(untagged)", 0)
        ws.cell(row=r, column=1, value=exam)
        ws.cell(row=r, column=2, value=subject)
        ws.cell(row=r, column=3, value=total)
        ws.cell(row=r, column=4, value=total - untagged)
        ws.cell(row=r, column=5, value=untagged)
        ws.cell(row=r, column=6, value=f"{round(100 * (total - untagged) / total)}%")
        ws.cell(row=r, column=7, value=len([t for t in counter if t != "(untagged)"]))
        r += 1
    for col, w in enumerate([7, 20, 12, 10, 10, 10, 15], start=1):
        ws.column_dimensions[get_column_letter(col)].width = w

    return wb


def main() -> int:
    if not INDEX.exists():
        print("FAIL: run scripts/build_bank.py first (no public/data/index.json)")
        return 1
    questions = load_bank()
    wb = build_workbook(questions)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Wrote {OUT} ({len(questions)} questions in bank, stratified sample inside)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
